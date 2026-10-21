"""Check dsh font selection against actual Latin/CJK PDF spans and warm reuse.

Requires python-docx, python-pptx, Poppler and the packaged macOS office/office-session diagnostic hosts.
Fixtures and measurements are generated in --output, never committed.
"""
import argparse
import html
import io
import json
from pathlib import Path
import selectors
import shutil
import statistics
import subprocess
import time
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from pptx import Presentation
from pptx.util import Inches


def document(path, family, text='Font ABC 中文测试', extra=None):
    family = html.escape(family, quote=True)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('[Content_Types].xml', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        archive.writestr('_rels/.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        archive.writestr('word/document.xml', f'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:rPr><w:rFonts w:ascii="{family}" w:hAnsi="{family}" w:eastAsia="{family}" w:cs="{family}"/></w:rPr><w:t>{html.escape(text)}</w:t></w:r></w:p></w:body></w:document>')
        for name, value in (extra or {}).items():
            archive.writestr(name, value)


def themed_document(path, family):
    # Use full package templates so themes are actually linked into the document.
    data = io.BytesIO()
    if path.suffix == '.docx':
        source = Document()
        source.add_paragraph('Theme sentinel')
        source.save(data)
        theme = 'word/theme/theme1.xml'
    else:
        source = Presentation()
        slide = source.slides.add_slide(source.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = 'Theme sentinel'
        source.save(data)
        theme = 'ppt/theme/theme1.xml'
    with zipfile.ZipFile(data) as original, zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as output:
        for name in original.namelist():
            content = original.read(name)
            if name == theme:
                xml = ET.fromstring(content)
                xml.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}minorFont/{http://schemas.openxmlformats.org/drawingml/2006/main}latin').set('typeface', family)
                content = ET.tostring(xml)
            output.writestr(name, content)


def spans(path):
    root = ET.fromstring(subprocess.check_output(['pdftohtml', '-xml', '-i', '-stdout', str(path)], timeout=20))
    result = []
    for page in root.findall('page'):
        fonts = {item.attrib['id']: item.attrib['family'].split('+')[-1] for item in page.findall('fontspec')}
        result += [(''.join(item.itertext()), fonts[item.attrib['font']]) for item in page.findall('text')]
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--office', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    office, root = args.office.resolve(), args.output.resolve()
    root.mkdir(parents=True, exist_ok=False)
    runtime = root/'runtime'; runtime.mkdir()
    log = (root/'session.log').open('w')
    session = subprocess.Popen([str(office.with_name('office-session')), str(office.parent.parent/'office/Frameworks'), str(runtime), '--timings'],
                               stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=log, text=True)
    records = []

    def convert(source, *options, success=True, output_format="pdf"):
        target = root/f'result-{len(records)}.{output_format}'
        start = time.perf_counter()
        session.stdin.write('\t'.join(map(str, ['convert', '--input', source, '--output', target, *options]))+'\n'); session.stdin.flush()
        poll = selectors.DefaultSelector(); poll.register(session.stdout, selectors.EVENT_READ)
        try: assert poll.select(90), 'font session timed out'
        finally: poll.close()
        result = json.loads(session.stdout.readline()); assert result['ok'] == success, result
        if not success: assert not target.exists()
        record = dict(ms=(time.perf_counter()-start)*1000, result=result, spans=spans(target) if success and output_format == "pdf" else [])
        records.append(record)
        return record

    try:
        # Exact reported failures, including localized aliases. These assertions
        # need PingFang/Songti installed on the validation Mac; do not silently skip.
        for family, latin, cjk, missing in [
            ('Microsoft YaHei', 'PingFangSC', 'PingFangSC', True),
            ('PingFang SC', 'PingFangSC', 'PingFangSC', False),
            ('SimSun', 'STSongti-SC', 'STSongti-SC', True),
            ('微软雅黑', 'PingFangSC', 'PingFangSC', True),
            ('宋体', 'STSongti-SC', 'STSongti-SC', True),
            ('Calibri', 'ArialMT', 'PingFangSC', True),
            ('Helvetica', 'Helvetica', 'PingFangSC', False),
        ]:
            source = root/f'{family}.docx'; document(source, family)
            result = convert(source)
            assert result['result']['missingFonts'] == ([family] if missing else []), result
            assert any('Font ABC' in text and font == latin for text, font in result['spans']), result
            assert any('中文测试' in text and font == cjk for text, font in result['spans']), result
        document(root/'cambria.docx', 'Cambria', 'Font ABC')
        result = convert(root/'cambria.docx')
        assert result['result']['missingFonts'] == ['Cambria'] and result['spans'][0][1] == 'TimesNewRomanPSMT', result

        # Metadata-only parts must not add defaults or themes to diagnostics.
        document(root/'helvetica.docx', 'Helvetica', 'Font ABC', {
            'word/fontTable.xml': '<w:fonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:rFonts w:ascii="UnusedTableFamily"/></w:fonts>',
            'word/theme/theme1.xml': '<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:latin typeface="UnusedThemeFamily"/></a:theme>',
        })
        result = convert(root/'helvetica.docx')
        assert result['result']['missingFonts'] == [] and result['spans'][0][1] == 'Helvetica', result
        # Installed theme fonts must be visible without an explicit declaration.
        # Alternate families on the same engine.
        theme_indices = []
        for kind in ['docx', 'pptx']:
            for family, expected in [('Georgia', 'Georgia'), ('Gill Sans', 'GillSans'), ('Georgia', 'Georgia')]:
                path = root/f'theme-{kind}-{len(records)}.{kind}'; themed_document(path, family)
                result = convert(path)
                assert result['result']['missingFonts'] == [], result
                assert any('Theme sentinel' in text and font == expected for text, font in result['spans']), result
                theme_indices.append(len(records)-1)
        # Installed originals win even if mentioned in a custom fallback group.
        result = convert(root/'helvetica.docx', '--font-fallbacks', '[["Helvetica","Courier New"]]')
        assert result['spans'][0][1] == 'Helvetica', result

        # A requested group member (not just its head) uses that group's order;
        # later generic/monospace membership must not overwrite it.
        document(root/'missing.docx', 'MissingOfficeFamily', 'Font ABC')
        both = root/'both'; both.mkdir()
        for font in ['Arial.ttf', 'Courier New.ttf']:
            shutil.copy('/System/Library/Fonts/Supplemental/'+font, both/font)
        for replacement, expected in [('Courier New', 'CourierNewPSMT'), ('Arial', 'ArialMT'), ('Courier New', 'CourierNewPSMT')]:
            result = convert(root/'missing.docx', '--font-directory', both, '--font-fallbacks', json.dumps([[replacement, 'MissingOfficeFamily'], ['monospace', 'MissingOfficeFamily', 'ArialMT']]))
            assert result['result']['missingFonts'] == ['MissingOfficeFamily'] and result['spans'][0][1] == expected, result

        # Catalog restriction must also remove previously loaded theme fonts,
        # and switching back must restore them on the same persistent engine.
        path = root/'theme-restricted.docx'; themed_document(path, 'Georgia')
        for options, expected in [([], 'Georgia'), (['--font-directory', both], 'ArialMT'), ([], 'Georgia')]:
            result = convert(path, *options)
            # Template styles also explicitly declare Courier; that independent
            # diagnostic is valid when the restricted catalog lacks it.
            assert 'Georgia' not in result['result']['missingFonts'], result
            assert any('Theme sentinel' in text and font == expected for text, font in result['spans']), result

        # The same search roots must notice file replacement and newly created
        # descendants. No engine restart or different --font-directory is used.
        mutable = root/'mutable'; mutable.mkdir()
        shutil.copy('/System/Library/Fonts/Supplemental/Arial.ttf', mutable/'font.ttf')
        options = ['--font-directory', mutable, '--font-fallbacks', '[["MissingOfficeFamily","Courier New","Arial"]]']
        result = convert(root/'missing.docx', *options); assert result['spans'][0][1] == 'ArialMT', result
        shutil.copy('/System/Library/Fonts/Supplemental/Courier New.ttf', mutable/'font.ttf')
        result = convert(root/'missing.docx', *options); assert result['spans'][0][1] == 'CourierNewPSMT', result
        shutil.copy('/System/Library/Fonts/Supplemental/Arial.ttf', mutable/'font.ttf')
        nested = mutable/'new'; nested.mkdir()
        shutil.copy('/System/Library/Fonts/Supplemental/Courier New.ttf', nested/'courier.ttf')
        result = convert(root/'missing.docx', *options); assert result['spans'][0][1] == 'CourierNewPSMT', result
        (nested/'courier.ttf').unlink()
        result = convert(root/'missing.docx', *options); assert result['spans'][0][1] == 'ArialMT', result

        # A deliberately declared absent engine-default name still gets a
        # diagnostic: this is document-based reporting, not a name blacklist.
        document(root/'liberation.docx', 'Liberation Sans', 'Font ABC')
        result = convert(root/'liberation.docx', '--font-directory', both)
        assert result['result']['missingFonts'] == ['Liberation Sans'], result
        convert(root/'helvetica.docx')  # Switch back to the default catalog.
        # Thousands of declarations exercise bounded native work on a warm
        # catalog. The extra XML part is inspected but not used by LibreOffice.
        for count in [4095, 4096, 5000]:
            xml = '<r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'+''.join(f'<w:rFonts w:ascii="Absent{i}"/>' for i in range(count))+'</r>'
            path = root/f'families-{count}.docx'
            document(path, 'Helvetica', 'Font ABC', {'word/extra.xml': xml})
            result = convert(path)
            assert len(result['result']['missingFonts']) == 4095, result
            assert result['result']['fontMetadataTruncated'] == (count > 4095), result
            many_family_index = len(records)-1
        # Actual body text (not an unused metadata part) must survive import and
        # export above the code-point sampling cap, including its final sentinel.
        cjk = ''.join(chr(p) for first,last in [(0x3400,0x4dc0),(0x4e00,0xa000),(0x20000,0x2a6e0)] for p in range(first,last))
        for count in [60000, 70000]:
            path = root/f'characters-{count}.docx'
            text = cjk[:count]+' End sentinel'
            document(path, 'Songti SC', text)
            result = convert(path, output_format='docx')
            assert result['result']['fontMetadataTruncated'] == (count > 65536), result
            with zipfile.ZipFile(result['result']['outputPath']) as output:
                xml = ET.fromstring(output.read('word/document.xml'))
            actual = ''.join(node.text or '' for node in xml.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
            assert actual == text, (count, len(actual), len(text))
        path = root/'long-family.docx'; document(path, 'X'*1025, 'Font ABC')
        result = convert(path)
        assert result['result']['fontMetadataTruncated'] and result['result']['missingFonts'] == [], result
        assert any('Font ABC' in text for text, font in result['spans']), result
        convert(root/'helvetica.docx')
        warm_start = len(records)
        for _ in range(16):
            result = convert(root/'helvetica.docx'); assert result['result']['missingFonts'] == [], result
    finally:
        session.stdin.close()
        try: assert session.wait(timeout=10) == 0
        finally: log.close()
    phases = [json.loads(line.removeprefix('OFFICE_PHASES ')) for line in (root/'session.log').read_text().splitlines() if line.startswith('OFFICE_PHASES ')]
    assert len(phases) == len(records)
    for index in theme_indices:
        assert not {"font-catalog-rebuilt", "fonts-reloaded"} & phases[index].keys(), phases[index]
    for phase in phases[warm_start:]:
        assert not {'font-catalog-rebuilt', 'fonts-reloaded', 'font-substitutions-updated'} & phase.keys(), phase
    # Cancellation after font activation must commit cache state before returning:
    # A -> cancelled B -> A must restore A, not believe the previous A is active.
    manifest = root/'cancel-fonts.tsv'
    manifest.write_text(''.join('\t'.join(map(str, [label, 'convert', root/source, root/target, '-', status, phase]))+'\n'
        for label, source, target, status, phase in [
            ('font-a', 'helvetica.docx', 'before-cancel.pdf', 0, '-'),
            ('font-b', 'SimSun.docx', 'cancelled.pdf', 130, 'font-substitutions-updated'),
            ('font-a-restored', 'helvetica.docx', 'after-cancel.pdf', 0, '-')]))
    probe = subprocess.run([str(office.with_name('office-probe')), '--memory-probe', str(manifest),
        str(root/'cancel-fonts.jsonl'), str(office.parent.parent/'office/Frameworks'), 'owner'], capture_output=True, text=True, timeout=90)
    assert probe.returncode == 0 and not (root/'cancelled.pdf').exists(), (probe.stdout, probe.stderr)
    assert spans(root/'after-cancel.pdf')[0][1] == 'Helvetica'
    report = {'passed': True, 'operations': len(records),
              'manyFamilySetupMs': phases[many_family_index]['fonts-configured']-phases[many_family_index]['runtime-ready'],
              'warmConversionMedianMs': statistics.median(r['ms'] for r in records[warm_start:]),
              'warmFontSetupMedianMs': statistics.median(p['fonts-configured']-p['runtime-ready'] for p in phases[warm_start:]),
              'records': records, 'phases': phases}
    (root/'results.json').write_text(json.dumps(report, indent=2)+'\n')
    print(json.dumps({key: value for key, value in report.items() if key not in ['records', 'phases']}))


if __name__ == '__main__':
    main()
