"""Generate and verify real Office fixtures against the standalone native shell.

Requires python-docx, python-pptx, openpyxl, pypdf, Pillow, and pdftoppm.
All artifacts are retained in a fresh caller-selected directory.
"""
import argparse
import hashlib
import json
import pathlib
import subprocess
import time

from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--office", type=pathlib.Path, required=True)
    parser.add_argument("--output", type=pathlib.Path, required=True)
    args = parser.parse_args()
    office = args.office.resolve(strict=True)
    root = args.output.resolve()
    root.mkdir(parents=True, exist_ok=False)
    records = []

    def run(*arguments, success=True):
        started = time.monotonic()
        process = subprocess.run(
            [str(office), *map(str, arguments)], capture_output=True, text=True,
            timeout=180,
        )
        result = json.loads(process.stdout if success else process.stderr.splitlines()[-1])
        records.append({"arguments": list(map(str, arguments)),
                        "seconds": time.monotonic() - started,
                        "exitCode": process.returncode, "result": result,
                        "stderr": process.stderr})
        assert (process.returncode == 0) == success, records[-1]
        assert result["ok"] == success, result
        assert not list(root.glob(".office-*")), "Private working directory leaked"
        return result

    document = Document()
    document.add_heading("Native office document", 0)
    document.add_paragraph("Document fixture: editable text, a table, and a page break.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Product"
    table.cell(0, 1).text = "Quantity"
    table.cell(1, 0).text = "Widget"
    table.cell(1, 1).text = "7"
    document.add_page_break()
    document.add_paragraph("Second page sentinel")
    document.save(root / "input document.docx")

    workbook = Workbook()
    inputs = workbook.active
    inputs.title = "Inputs"
    inputs.append(["Quantity", "Price"])
    inputs.append([3, 19])
    inputs.append([4, 12])
    summary = workbook.create_sheet("Summary")
    summary.append(["Native office workbook", "Value"])
    summary.append(["Revenue", "=Inputs!A2*Inputs!B2+Inputs!A3*Inputs!B3"])
    summary.append(["Double revenue", "=B2*2"])
    summary.column_dimensions["A"].width = 32
    workbook.save(root / "input workbook.xlsx")

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = "Native office presentation"
    presentation.save(root / "input slides.pptx")
    input_hashes = {path.name: hashlib.sha256(path.read_bytes()).hexdigest()
                    for path in root.glob("input *")}

    pdf_reports = []
    for filename, phrase, count in [
        ("input document.docx", "Second page sentinel", 2),
        ("input workbook.xlsx", "Native office workbook", None),
        ("input slides.pptx", "Native office presentation", 1),
    ]:
        source = root / filename
        output = source.with_suffix(".pdf")
        run("convert", "--input", source, "--output", output)
        pdf = PdfReader(output)
        text = "\n".join(page.extract_text() for page in pdf.pages)
        assert phrase in text, (filename, text)
        if count is not None:
            assert len(pdf.pages) == count, (filename, len(pdf.pages))
        assert output.stat().st_mode & 0o777 == 0o600
        prefix = output.with_suffix("")
        subprocess.run(["pdftoppm", "-f", "1", "-singlefile", "-r", "72", "-png",
                        str(output), str(prefix)], check=True, capture_output=True, timeout=60)
        with Image.open(prefix.with_suffix(".png")) as image:
            dark_pixels = sum(value < 200 for value in image.convert("L").getdata())
            assert dark_pixels > 100, (filename, dark_pixels)
            pdf_reports.append({"input": filename, "pages": len(pdf.pages),
                                "firstPageSize": image.size, "darkPixels": dark_pixels})

    calculated = root / "calculated.xlsx"
    run("recalculate", "--input", root / "input workbook.xlsx", "--output", calculated)
    formulas = load_workbook(calculated, data_only=False)
    values = load_workbook(calculated, data_only=True)
    assert formulas["Summary"]["B2"].data_type == "f"
    assert formulas["Summary"]["B3"].data_type == "f"
    assert values["Summary"]["B2"].value == 105
    assert values["Summary"]["B3"].value == 210
    run("convert", "--input", calculated, "--output", root / "summary.csv", "--sheet", "Summary")
    assert "105" in (root / "summary.csv").read_text()
    # An existing output and invalid document must not destroy prior artifacts.
    prior = hashlib.sha256(calculated.read_bytes()).hexdigest()
    run("recalculate", "--input", root / "input workbook.xlsx", "--output", calculated, success=False)
    assert hashlib.sha256(calculated.read_bytes()).hexdigest() == prior
    invalid = root / "invalid.docx"
    invalid.write_text("This is not an Office ZIP container.")
    run("convert", "--input", invalid, "--output", root / "invalid.pdf", success=False)
    assert not (root / "invalid.pdf").exists()
    run("convert", "--input", calculated, "--output", root / "missing-sheet.csv",
        "--sheet", "NotASheet", success=False)
    assert not (root / "missing-sheet.csv").exists()
    for name, expected in input_hashes.items():
        assert hashlib.sha256((root / name).read_bytes()).hexdigest() == expected
    report = {"passed": True, "pdfs": pdf_reports,
              "formulaCaches": {"Summary!B2": 105, "Summary!B3": 210},
              "commands": records}
    (root / "report.json").write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps({"passed": True, "commands": len(records), "report": str(root / "report.json")}))


if __name__ == "__main__":
    main()
