"""Exercise rendering, fonts, container guards and the public command contract.

Run with the same Python dependencies as verify_runtime.py; no Node host is used.
"""
import argparse
import hashlib
import io
import json
import os
from pathlib import Path
import selectors
import shutil
import struct
import subprocess
import zipfile
import zlib

from docx import Document
from openpyxl import load_workbook
from PIL import Image


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--office', type=Path, required=True)
    parser.add_argument('--fixtures', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    office, fixtures, root = args.office.resolve(), args.fixtures.resolve(), args.output.resolve()
    root.mkdir(parents=True, exist_ok=False)
    records = []

    def run(*words, success=True):
        p = subprocess.run([str(office), *map(str, words)], text=True, capture_output=True, timeout=90)
        assert p.returncode == (0 if success else 1), (words, p.stdout, p.stderr)
        value = json.loads(p.stdout if success else p.stderr.splitlines()[-1])
        assert value['ok'] == success, value
        if success: assert value['fontMetadataTruncated'] is False, value
        if not success:
            assert not p.stdout and value['code'] and value['error'], p
        records.append(dict(arguments=list(map(str, words)), result=value))
        return value

    for words in [[], ['--help'], ['-h'], ['help'], ['convert', '--help'], ['render', '-h'], ['recalculate', '--help']]:
        p = subprocess.run([str(office), *words], capture_output=True, text=True, timeout=5, cwd='/')
        assert p.returncode == 0 and '600' in p.stdout and '130' in p.stdout and not p.stderr
        assert '--initial-font-family' not in p.stdout
    assert subprocess.check_output([str(office), '--version'], text=True).startswith('office ')
    caps = json.loads(subprocess.check_output([str(office), 'capabilities', '--json'], text=True))
    assert json.loads(subprocess.check_output([str(office), 'capabilities'], text=True)) == caps
    assert caps['imageRendering']['pdfBackend'] == 'pdfium' and not caps['execution']['hardTimeout']
    assert caps['fonts']['allCatalogFontsAvailable'] and caps['fonts']['metadataSampling']['maxFamilies'] == 4096
    assert caps['fonts']['metadataSampling']['maxCodePoints'] == 65536 and caps['fonts']['metadataSampling']['maxFamilyBytes'] == 1024
    assert caps['fonts']['metadataSampling']['overflow'] == 'truncate'
    assert caps['fonts']['options'] == ['font-directory', 'font-fallbacks']

    def pngs(result, blank=False):
        for item in result['images']:
            path = Path(item['path'])
            raw = path.read_bytes()
            assert len(raw) == item['byteLength'] and raw[:8] == b'\x89PNG\r\n\x1a\n'
            cursor = 8
            while cursor < len(raw):
                size = struct.unpack_from('>I', raw, cursor)[0]
                chunk = raw[cursor+4:cursor+8+size]
                assert zlib.crc32(chunk) == struct.unpack_from('>I', raw, cursor+8+size)[0]
                cursor += size + 12
            assert cursor == len(raw)
            with Image.open(path) as image:
                image.load()
                assert image.size == (item['width'], item['height'])
                if not blank:
                    assert any(lo != hi for lo, hi in image.convert('RGB').getextrema()), path

    for index, name in enumerate(['input document.docx', 'input slides.pptx', 'input workbook.xlsx']):
        result = run('render', '--input', fixtures/name, '--output-dir', root/f'render-{index}')
        assert result['sourceSha256'] == hashlib.sha256((fixtures/name).read_bytes()).hexdigest()
        pngs(result)
    result = run('render', '--input', fixtures/'input document.docx', '--output-dir', root/'selected', '--pages', '2,1', '--dpi', '96')
    assert [p['page'] for p in result['images']] == [2, 1]
    assert result['images'][0]['width'] == 816
    pngs(result)
    run('convert', '--input', fixtures/'input document.docx', '--output', root/'input.pdf')
    result = run('render', '--input', root/'input.pdf', '--output-dir', root/'pdf')
    assert result['rasterEngine'] == 'pdfium' and result['pageCount'] == 2
    pngs(result)
    wb = load_workbook(fixtures/'input workbook.xlsx')
    wb['Summary'].sheet_view.rightToLeft = True
    wb['Summary'].column_dimensions['B'].hidden = True
    hidden = wb.create_sheet('Hidden'); hidden['A1'] = 'Hidden sentinel'; hidden.sheet_state = 'hidden'
    wb.save(root/'sheets.xlsx')
    result = run('render', '--input', root/'sheets.xlsx', '--output-dir', root/'sheets')
    assert result['pageCount'] == 2 and len(result['images']) == 2
    assert result['images'][1]['rtl']
    pngs(result)
    result = run('render', '--input', root/'sheets.xlsx', '--output-dir', root/'hidden', '--sheet', 'Hidden', '--range', '$A$1:D20', '--max-pixels', '10000', '--max-dimension', '100')
    assert len(result['images']) > 1 and all(i['width'] * i['height'] <= 10000 for i in result['images'])
    pngs(result, blank=True)
    for index, options in enumerate([
        ['--max-pages', '1'], ['--pages', '1,1'], ['--pages', '3'], ['--pages', '1,'],
        ['--dpi', '601'], ['--max-pixels', '10'], ['--max-output-bytes', '30'],
        ['--max-input-bytes', '1'], ['--max-archive-entries', '1'], ['--max-uncompressed-bytes', '1'],
        ['--max-output-bytes', '99999999999999999999999'], ['--dpi', '-1'], ['--dpi', '96', '--dpi', '144'],
        ['--initial-font-family', 'Helvetica'], ['--font-fallbacks', 'bad'], ['--font-fallbacks', '[["single"]]'], ['--font-directory', str(root/'missing')],
    ]):
        target = root/f'failed-{index}'
        result = run('render', '--input', fixtures/'input document.docx', '--output-dir', target, *options, success=False)
        assert 'stoull' not in result['error'] and not target.exists()
    empty = root/'empty.docx'; empty.touch()
    directory = root/'directory.docx'; directory.mkdir()
    for source in [empty, directory]:
        failure = run('convert', '--input', source, '--output', root/(source.stem+'.pdf'), success=False)
        assert failure['code'] == 'invalid-document', failure
    original = (root/'selected/page-0001.png').read_bytes()
    run('render', '--input', fixtures/'input document.docx', '--output-dir', root/'selected', success=False)
    assert (root/'selected/page-0001.png').read_bytes() == original
    run('render', '--input', root/'sheets.xlsx', '--output-dir', root/'wrong-sheet', '--sheet', 'summary', success=False)
    run('render', '--input', root/'sheets.xlsx', '--output-dir', root/'no-sheet', '--range', 'A1:B2', success=False)

    # Accept real streamed ZIP descriptors and ZIP64 local headers, not only
    # the ordinary seekable ZIPs produced by the fixture libraries.
    class Stream(io.BytesIO):
        def seekable(self): return False
        def seek(self, *args): raise io.UnsupportedOperation('stream')
    for wide in [False, True]:
        stream = Stream()
        with zipfile.ZipFile(fixtures/'input document.docx') as original, zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as archive:
            for name in original.namelist():
                with archive.open(name, 'w', force_zip64=wide) as output:
                    output.write(original.read(name))
        path = root/f'stream-{wide}.docx'; path.write_bytes(stream.getvalue())
        run('convert', '--input', path, '--output', root/f'stream-{wide}.pdf')

    source = (fixtures/'input document.docx').read_bytes()
    # The preflight accepts upstream's permissive ZIP policy; exercise changes
    # that LibreOffice itself also accepts through the complete conversion path.
    tolerated = {'trailing-data': source+b'trailing data'}
    for name, data in tolerated.items():
        path = root/(name+'.docx'); path.write_bytes(data)
        run('convert', '--input', path, '--output', root/(name+'.pdf'))
    malformed = {'truncated': source[:-12], 'wrong-kind': (fixtures/'input workbook.xlsx').read_bytes()}
    corrupt = bytearray(source)
    corrupt[30+struct.unpack_from('<H', corrupt, 26)[0]+struct.unpack_from('<H', corrupt, 28)[0]+2] ^= 0xff
    malformed['bad-deflate'] = corrupt
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('large', b'X' * 2**20)
    malformed['bomb'] = bomb.getvalue()
    for name, data in malformed.items():
        file = root/(name+'.docx'); file.write_bytes(data)
        run('convert', '--input', file, '--output', root/(name+'.pdf'), '--max-uncompressed-bytes', '262144', success=False)
    ole = root/'bad.doc'; ole.write_bytes(bytes.fromhex('d0cf11e0a1b11ae1')+bytes(1016))
    run('convert', '--input', ole, '--output', root/'ole.pdf', success=False)

    # A/B/A on the same owner proves settings are refreshed, not just parsed.
    for name, files in [('a', ['Arial.ttf']), ('b', ['Courier New.ttf']), ('both', ['Arial.ttf', 'Courier New.ttf'])]:
        directory = root/name; directory.mkdir()
        for file in files: shutil.copy('/System/Library/Fonts/Supplemental/'+file, directory/file)
    document = Document()
    document.add_paragraph().add_run('Font metrics www iii 12345').font.name = 'MissingOfficeSentinel'
    document.save(root/'fonts.docx')
    runtime = root/'runtime'; runtime.mkdir()
    log = (root/'session.log').open('w')
    session = subprocess.Popen([str(office.with_name('office-session')), str(office.parent.parent/'office/Frameworks'), str(runtime)],
                               stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=log, text=True)
    def call(words):
        session.stdin.write('\t'.join(map(str, words))+'\n'); session.stdin.flush()
        poll = selectors.DefaultSelector(); poll.register(session.stdout, selectors.EVENT_READ)
        try: assert poll.select(60), 'session timeout'
        finally: poll.close()
        result = json.loads(session.stdout.readline()); assert result['ok'], result
        return result
    expected = [('a','Arial','ArialMT'), ('b','Courier New','CourierNewPSMT'), ('a','Arial','ArialMT'),
                ('both','Courier New','CourierNewPSMT'), ('both','Arial','ArialMT')]
    try:
        for index, (directory, family, embedded) in enumerate(expected):
            target = root/f'font-{index}.pdf'
            result = call(['convert', '--input', root/'fonts.docx', '--output', target, '--font-directory', root/directory,
                          '--font-fallbacks', json.dumps([['MissingOfficeSentinel',family]])])
            assert 'MissingOfficeSentinel' in result['missingFonts']
            fonts = subprocess.check_output(['pdffonts', str(target)], text=True)
            assert embedded in fonts, (index, fonts)
        active = list(runtime.glob('.office-runtime-*')); assert len(active) == 1
        other = subprocess.run([str(office.with_name('office-session')), str(office.parent.parent/'office/Frameworks'), str(runtime)],
                               input='\t'.join(map(str,['convert','--input',root/'fonts.docx','--output',root/'concurrent.pdf']))+'\n',
                               capture_output=True, text=True, timeout=60)
        assert other.returncode == 0 and json.loads(other.stdout)['ok'], other.stderr
        assert active[0].is_dir() and len(list(runtime.glob('.office-runtime-*'))) == 2
    finally:
        session.stdin.close(); assert session.wait(timeout=10) == 0; log.close()
    # A later process removes its predecessor's leased profile, preserving the cache.
    p = subprocess.run([str(office.with_name('office-session')), str(office.parent.parent/'office/Frameworks'), str(runtime)],
                       input='\t'.join(map(str,['convert','--input',root/'fonts.docx','--output',root/'next.pdf']))+'\n',
                       capture_output=True, text=True, timeout=60)
    assert p.returncode == 0 and json.loads(p.stdout)['ok'], p.stderr
    assert not active[0].exists() and len(list(runtime.glob('.office-runtime-*'))) == 1
    assert (runtime/'font-cache').is_dir()
    assert not list(root.glob('.office-*')), 'per-document scratch leaked'
    (root/'results.json').write_text(json.dumps(records, indent=2)+'\n')
    print(json.dumps({'passed': True, 'commands': len(records), 'fontTransitions': len(expected)}))

if __name__ == '__main__':
    main()
