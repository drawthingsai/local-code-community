#!/usr/bin/env python3
"""Reopen an OOXML workbook and report compact structural checks."""

from __future__ import annotations

import argparse
import json
import zipfile
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple


ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")


def project_file(value: str, root_value: str) -> tuple[Path, Path]:
    supplied = Path(value)
    if supplied.is_absolute() or ".." in supplied.parts:
        raise ValueError("Workbook path must be project-relative and may not contain '..'.")
    root = Path(root_value).resolve()
    if not root.is_dir():
        raise ValueError(f"Project root is not a directory: {root_value}")
    resolved = (root / supplied).resolve()
    try:
        resolved.relative_to(root)
    except ValueError as error:
        raise ValueError("Workbook path resolves outside the selected project.") from error
    if not resolved.is_file():
        raise ValueError(f"Workbook does not exist: {value}")
    if resolved.suffix.lower() not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        raise ValueError("Expected an OOXML workbook (.xlsx, .xlsm, .xltx, or .xltm).")
    return root, resolved


def json_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return repr(value)


def parse_cell_reference(reference: str) -> tuple[str, str]:
    if "!" not in reference:
        raise ValueError(f"Cell reference must use Sheet!A1 form: {reference}")
    sheet_name, coordinate = reference.rsplit("!", 1)
    sheet_name = sheet_name.strip()
    if len(sheet_name) >= 2 and sheet_name[0] == sheet_name[-1] == "'":
        sheet_name = sheet_name[1:-1].replace("''", "'")
    coordinate = coordinate.strip().upper()
    try:
        coordinate_to_tuple(coordinate)
    except ValueError as error:
        raise ValueError(f"Invalid cell coordinate in {reference}") from error
    if not sheet_name:
        raise ValueError(f"Missing sheet name in {reference}")
    return sheet_name, coordinate


def nested_value(value: Any, *names: str) -> Any:
    for name in names:
        value = getattr(value, name, None)
        if value is None:
            return None
    return value


def chart_report(chart: Any) -> dict[str, Any]:
    references: set[str] = set()
    paths = (
        ("tx", "strRef", "f"),
        ("cat", "strRef", "f"),
        ("cat", "numRef", "f"),
        ("val", "numRef", "f"),
        ("xVal", "numRef", "f"),
        ("yVal", "numRef", "f"),
    )
    for series in chart.ser:
        for path in paths:
            reference = nested_value(series, *path)
            if isinstance(reference, str):
                references.add(reference)
    anchor = chart.anchor
    if isinstance(anchor, str):
        anchor_text = anchor
    elif getattr(anchor, "_from", None) is not None:
        marker = anchor._from
        anchor_text = f"{get_column_letter(marker.col + 1)}{marker.row + 1}"
    else:
        anchor_text = None
    return {
        "anchor": anchor_text,
        "series_references": sorted(references),
        "type": type(chart).__name__,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", help="Project-relative workbook path")
    parser.add_argument(
        "--project-root", default=".", metavar="PATH",
        help="Project root used to resolve and contain the workbook path",
    )
    parser.add_argument(
        "--cell", action="append", default=[], metavar="SHEET!A1",
        help="Representative cell to report; may be repeated",
    )
    parser.add_argument(
        "--expect-sheet", action="append", default=[], metavar="NAME",
        help="Sheet that must exist; may be repeated",
    )
    args = parser.parse_args()

    try:
        root, path = project_file(args.workbook, args.project_root)
        with zipfile.ZipFile(path) as archive:
            corrupt_member = archive.testzip()
        if corrupt_member is not None:
            raise ValueError(f"Workbook ZIP contains a corrupt member: {corrupt_member}")

        keep_vba = path.suffix.lower() in {".xlsm", ".xltm"}
        formulas_book = load_workbook(path, data_only=False, keep_vba=keep_vba)
        values_book = load_workbook(path, data_only=True, read_only=True, keep_vba=keep_vba)
        issues: list[dict[str, str]] = []
        sheet_reports: list[dict[str, Any]] = []
        total_formulas = 0
        for sheet in formulas_book.worksheets:
            formula_count = 0
            nonempty_count = 0
            styled_count = 0
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if cell.has_style:
                        styled_count += 1
                    if value is None:
                        continue
                    nonempty_count += 1
                    upper = value.upper() if isinstance(value, str) else ""
                    if isinstance(value, str) and value.startswith("="):
                        formula_count += 1
                        for token in ERROR_TOKENS:
                            if token in upper:
                                issues.append(
                                    {"cell": f"{sheet.title}!{cell.coordinate}", "issue": f"Formula contains {token}"}
                                )
                        cached_value = values_book[sheet.title][cell.coordinate].value
                        if isinstance(cached_value, str) and cached_value.upper() in ERROR_TOKENS:
                            issues.append(
                                {"cell": f"{sheet.title}!{cell.coordinate}", "issue": f"Cached value contains {cached_value.upper()}"}
                            )
                    elif upper in ERROR_TOKENS:
                        issues.append(
                            {"cell": f"{sheet.title}!{cell.coordinate}", "issue": f"Cell contains {upper}"}
                        )
            total_formulas += formula_count
            sheet_reports.append(
                {
                    "charts": [chart_report(chart) for chart in sheet._charts],
                    "conditional_formatting": [str(item.sqref) for item in sheet.conditional_formatting],
                    "data_validations": [str(item.sqref) for item in sheet.data_validations.dataValidation],
                    "dimension": sheet.calculate_dimension(),
                    "formulas": formula_count,
                    "freeze_panes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
                    "max_column": sheet.max_column,
                    "max_row": sheet.max_row,
                    "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
                    "nonempty_cells": nonempty_count,
                    "sheet": sheet.title,
                    "styled_cells": styled_count,
                    "tables": [
                        {"display_name": table.displayName, "name": table.name, "ref": table.ref}
                        for table in sheet.tables.values()
                    ],
                }
            )

        for expected in args.expect_sheet:
            if expected not in formulas_book.sheetnames:
                issues.append({"sheet": expected, "issue": "Expected sheet is missing"})

        cells: list[dict[str, Any]] = []
        for reference in args.cell:
            sheet_name, coordinate = parse_cell_reference(reference)
            if sheet_name not in formulas_book.sheetnames:
                issues.append({"cell": reference, "issue": "Representative cell sheet is missing"})
                continue
            formula_cell = formulas_book[sheet_name][coordinate]
            value_cell = values_book[sheet_name][coordinate]
            cells.append(
                {
                    "cached_value": json_value(value_cell.value),
                    "cell": f"{sheet_name}!{coordinate}",
                    "data_type": formula_cell.data_type,
                    "formula_or_value": json_value(formula_cell.value),
                    "number_format": formula_cell.number_format,
                    "style_id": formula_cell.style_id,
                }
            )

        sheet_names = formulas_book.sheetnames
        formulas_book.close()
        values_book.close()
        result = {
            "bytes": path.stat().st_size,
            "formula_count": total_formulas,
            "issues": issues,
            "ok": not issues,
            "path": path.relative_to(root).as_posix(),
            "representative_cells": cells,
            "sheet_names": sheet_names,
            "sheets": sheet_reports,
            "visual_fidelity_checked": False,
            "zip_integrity": "ok",
        }
        print(json.dumps(result, indent=2, sort_keys=True))
        return 0 if not issues else 2
    except (OSError, ValueError, zipfile.BadZipFile) as error:
        print(json.dumps({"error": str(error), "ok": False}, sort_keys=True))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
